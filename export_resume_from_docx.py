"""
Extract plain text from a Word resume (.docx).

Usage:
  python export_resume_from_docx.py
  python export_resume_from_docx.py --docx path/to/file.docx --txt out.txt
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    import docx  # python-docx
except ImportError:
    print("Missing dependency: pip install python-docx", file=sys.stderr)
    sys.exit(1)


def normalize_line(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def extract_paragraphs(document: docx.document.Document) -> list[tuple[str, str]]:
    """Return list of (style_name, text) for each non-empty paragraph."""
    rows: list[tuple[str, str]] = []
    for para in document.paragraphs:
        style = (para.style.name if para.style else "") or ""
        text = normalize_line(para.text)
        if text:
            rows.append((style, text))
    return rows


def extract_tables(document: docx.document.Document) -> list[str]:
    lines: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [normalize_line(c.text) for c in row.cells]
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                lines.append(row_text)
    return lines


def build_plain_text(paragraphs: list[tuple[str, str]], tables: list[str]) -> str:
    parts: list[str] = []
    for _style, text in paragraphs:
        parts.append(text)
    if tables:
        parts.append("")
        parts.append("--- Tables ---")
        parts.extend(tables)
    return "\n\n".join(parts).strip() + "\n"


def parse_args() -> argparse.Namespace:
    root = Path(__file__).resolve().parent
    p = argparse.ArgumentParser(description="Export .docx resume to UTF-8 plain text.")
    p.add_argument(
        "--docx",
        type=Path,
        default=root / "Jasen_Lee.docx",
        help="Input Word file",
    )
    p.add_argument(
        "--txt",
        type=Path,
        default=root / "Jasen_Lee_extracted.txt",
        help="Output UTF-8 plain text",
    )
    return p.parse_args()


def main() -> int:
    args = parse_args()
    docx_path = args.docx.resolve()
    if not docx_path.is_file():
        print(f"File not found: {docx_path}", file=sys.stderr)
        return 1

    document = docx.Document(str(docx_path))
    paragraphs = extract_paragraphs(document)
    tables = extract_tables(document)
    plain = build_plain_text(paragraphs, tables)

    txt_path = args.txt.resolve()
    txt_path.parent.mkdir(parents=True, exist_ok=True)
    txt_path.write_text(plain, encoding="utf-8")
    print(f"Wrote text: {txt_path} ({len(plain)} chars, {len(paragraphs)} paragraphs)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
